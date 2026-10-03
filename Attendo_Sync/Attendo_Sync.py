"""Attendo-Sync: local attendance dashboard with optional Google Sheets backup."""

from __future__ import annotations

import argparse
import json
import os
import sys
import tempfile
import traceback
from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace
from tkinter import messagebox

import gspread
import ttkbootstrap
from zk import ZK

from attendance_data import AttendanceStore, month_display
from dashboard import AttendoSyncApp

# Next to the EXE when frozen, so the backup export lands beside it.
ROOT_DIR = Path(sys.executable).parent if getattr(sys, "frozen", False) else Path(__file__).resolve().parent
BUNDLE_DIR = Path(getattr(sys, "_MEIPASS", ROOT_DIR))
OUTPUT_FILE = ROOT_DIR / "attendance_export.json"
SAMPLE_FILE_NAME = "sample_essl_attendance.json"
ICON_ICO = "Atteno_Sync_Icon.ico"
ICON_PNG = "Atteno_Sync_Icon.png"

# Fixed values for this EXE.
DEFAULT_DEVICE_IP = "192.168.29.201"
DEFAULT_PORT = 4370
DEVICE_TIMEOUT = 10
DEFAULT_CREDENTIALS_FILE = "C:/secure/attendance/gcloud/attendo-sync-f3615fbacce2.json"
SPREADSHEET_ID = "1VbfEv-e-AEeYGobS-E8O31Zi1OYwtsPxLIXXXtAB0BQ"
WORKSHEET_NAME = "Eurobia Attendance"
SAMPLE_WORKSHEET_NAME = "Sample_Test"
BUILD_ID = "2026-10-01 dashboard"

BLUE = "#2080A6"
DARK = "#16607D"
BG = "#F3F6F8"
TEXT = "#1E2A32"
MUTED = "#5B6B76"


def fetch_raw_attendance(device_ip, port=4370, timeout=10):
    zk = ZK(device_ip, port=port, timeout=timeout, password=0, force_udp=False, ommit_ping=False)
    conn = zk.connect()
    try:
        users = conn.get_users()
        user_map = {u.user_id: u.name for u in users}
        attendance = conn.get_attendance()
        return attendance, user_map
    finally:
        conn.disconnect()


def attendance_to_rows(attendance, user_map):
    rows = []
    for record in attendance:
        user_id = getattr(record, "user_id", None)
        timestamp = getattr(record, "timestamp", None)
        rows.append({
            "user_id": user_id,
            "user_name": user_map.get(user_id, "Unknown"),
            "timestamp": timestamp.isoformat() if timestamp else None,
            "status": getattr(record, "status", None),
            "punch": getattr(record, "punch", None),
        })
    return rows


def fetch_records(device_ip: str, port: int):
    attendance, user_map = fetch_raw_attendance(device_ip, port=port, timeout=DEVICE_TIMEOUT)
    return attendance_to_rows(attendance, user_map)


def load_sample_records():
    """Loads eSSL-shaped sample punches and runs them through the same conversion as real device data."""
    for folder in (ROOT_DIR, BUNDLE_DIR):
        sample_path = folder / SAMPLE_FILE_NAME
        if sample_path.exists():
            break
    else:
        raise FileNotFoundError(f"{SAMPLE_FILE_NAME} not found next to the application.")

    with sample_path.open("r", encoding="utf-8") as fh:
        sample = json.load(fh)

    user_map = {u["user_id"]: u["name"] for u in sample["users"]}
    attendance = [
        SimpleNamespace(
            uid=r["uid"],
            user_id=r["user_id"],
            timestamp=datetime.strptime(r["timestamp"], "%Y-%m-%d %H:%M:%S"),
            status=r["status"],
            punch=r["punch"],
        )
        for r in sample["attendance"]
    ]
    return attendance_to_rows(attendance, user_map)


def export_rows(rows, output_path: Path):
    with output_path.open("w", encoding="utf-8") as fh:
        json.dump(rows, fh, indent=2)


def build_worksheet_rows(rows):
    uploaded_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    sheet_rows = [["user_id", "user_name", "timestamp", "status", "punch", "uploaded_at"]]
    for row in rows:
        sheet_rows.append([
            row.get("user_id"),
            row.get("user_name"),
            row.get("timestamp"),
            row.get("status"),
            row.get("punch"),
            uploaded_at,
        ])
    return sheet_rows


def upload_to_google_sheets(rows, credentials_path, worksheet_name, log=print, spreadsheet_id=SPREADSHEET_ID):
    credentials_file = Path(credentials_path)
    if not credentials_path or not credentials_file.is_file():
        raise ValueError(f"Key file not found: {credentials_path or '(empty)'}")

    log(f"Reading key file: {credentials_file}")
    try:
        key = json.loads(credentials_file.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        raise ValueError("The selected file is not valid JSON.") from None
    if not isinstance(key, dict):
        raise ValueError("The selected JSON is not an object, so it cannot be a service account key.")
    log("Fields found in file: " + ", ".join(sorted(key)))

    missing = [f for f in ("type", "client_email", "private_key", "token_uri") if f not in key]
    if missing or key.get("type") != "service_account":
        raise ValueError(
            "This is not a Google service account key (missing: " + ", ".join(missing or ["type=service_account"]) + ").\n"
            "Create one in Google Cloud Console: IAM & Admin > Service Accounts > your account > "
            "Keys > Add key > Create new key > JSON, then select that downloaded file."
        )
    log(f"Service account: {key['client_email']}")

    log("Connecting to Google Sheets...")
    client = gspread.service_account(filename=str(credentials_file))
    client.set_timeout(30)
    spreadsheet = client.open_by_key(spreadsheet_id)
    log(f"Opened spreadsheet: {spreadsheet.title}")

    try:
        worksheet = spreadsheet.worksheet(worksheet_name)
    except gspread.WorksheetNotFound:
        log(f"Tab '{worksheet_name}' not found, creating it.")
        worksheet = spreadsheet.add_worksheet(
            title=worksheet_name,
            rows=str(max(100, len(rows) + 20)),
            cols="20",
        )

    existing = worksheet.get_all_values()
    if not any(str(cell).strip() for row in existing for cell in row):
        existing = []
    header = build_worksheet_rows([])[0]

    def normalize(cell):
        return str(cell or "").strip().casefold().replace(" ", "_").replace("-", "_")

    if existing and [normalize(cell) for cell in existing[0][:5]] != header[:5]:
        found = ", ".join(str(cell) for cell in existing[0] if str(cell).strip()) or "(blank first row)"
        raise ValueError(
            f"The backup worksheet '{worksheet_name}' has an incompatible header.\n"
            f"Found in row 1: {found}\n"
            f"Expected: {', '.join(header)}\n"
            "In Settings, type a new worksheet name (the app creates it automatically), "
            "or clear the existing tab and retry."
        )

    def identity(values):
        padded = list(values) + [""] * max(0, 6 - len(values))
        return tuple("" if padded[index] is None else str(padded[index]) for index in (0, 2, 3, 4))

    known = {identity(row) for row in existing[1:]}
    missing_rows = [] if existing else [header]
    for row in build_worksheet_rows(rows)[1:]:
        key = identity(row)
        if key not in known:
            missing_rows.append(row)
            known.add(key)
    log(f"Backing up {len(missing_rows) - (0 if existing else 1)} new punches to '{worksheet_name}'...")
    for offset in range(0, len(missing_rows), 1000):
        worksheet.append_rows(missing_rows[offset:offset + 1000], value_input_option="RAW")
    return worksheet.url


def find_resource(name):
    for folder in (BUNDLE_DIR, ROOT_DIR):
        path = folder / name
        if path.exists():
            return path
    return None


def fetch_dashboard_records(device_ip, port):
    attendance, users = fetch_raw_attendance(device_ip, port=port, timeout=DEVICE_TIMEOUT)
    return attendance_to_rows(attendance, users), users


def standalone_self_test(report_path):
    from attendance_reports import export_workbook
    from openpyxl import load_workbook

    root = None
    result = {"frozen": bool(getattr(sys, "frozen", False)), "ok": False}
    try:
        with tempfile.TemporaryDirectory(prefix="attendo-self-test-") as folder:
            store = AttendanceStore(Path(folder) / "history.sqlite3")
            records = load_sample_records()
            store.merge(records)
            employees, rows = store.snapshot()
            if not employees or not rows:
                raise RuntimeError("Bundled sample records are empty.")
            root = ttkbootstrap.Window(themename="flatly")
            defaults = {"device_ip": DEFAULT_DEVICE_IP, "port": str(DEFAULT_PORT),
                        "credentials": "", "spreadsheet_id": "", "worksheet": ""}
            view = AttendoSyncApp(root, store, None, None, defaults, demo=True,
                                 logo_path=find_resource(ICON_PNG))
            month = rows[0]["timestamp"][:7]
            view.month.set(month_display(month))
            view.refresh_month()
            root.update()
            view.canvas.draw()
            view.figure.savefig(Path(folder) / "chart.png")
            export_workbook(Path(folder) / "report.xlsx", employees, rows, month)
            workbook = load_workbook(Path(folder) / "report.xlsx")
            try:
                if len(workbook.sheetnames) != 5 or len(view.notebook.tabs()) != 3:
                    raise RuntimeError("Dashboard or workbook is incomplete.")
            finally:
                workbook.close()
            view.close()
            root = None
            result.update(ok=True, employees=len(employees), punches=len(rows),
                          checks=["SQLite", "bundled sample", "Tk GUI", "chart rendering", "Excel export", "Google and device imports"])
    except Exception:
        result["error"] = traceback.format_exc()
    finally:
        if root is not None:
            root.destroy()
    report_path.write_text(json.dumps(result, indent=2), encoding="utf-8")
    return 0 if result["ok"] else 1


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--self-test", type=Path, metavar="REPORT_JSON", help="Test bundled resources offline and exit")
    parser.add_argument("--demo", action="store_true", help="Open isolated sample data without device or cloud access")
    parser.add_argument("--data-dir", type=Path, default=Path(os.getenv("LOCALAPPDATA", str(ROOT_DIR))) / "AttendoSync")
    args = parser.parse_args()
    if args.self_test:
        raise SystemExit(standalone_self_test(args.self_test))
    root = ttkbootstrap.Window(themename="flatly")
    root.withdraw()
    try:
        args.data_dir.mkdir(parents=True, exist_ok=True)
        store = AttendanceStore(args.data_dir / ("demo.sqlite3" if args.demo else "attendance.sqlite3"))
        if args.demo and not store.snapshot()[1]:
            rows = load_sample_records()
            last_day = max(datetime.fromisoformat(row["timestamp"]).date() for row in rows)
            shift = date.today() - last_day
            for row in rows:
                row["timestamp"] = (datetime.fromisoformat(row["timestamp"]) + shift).isoformat()
            store.merge(rows)
        if not args.demo and not store.setting("legacy_imported") and OUTPUT_FILE.is_file():
            store.merge(json.loads(OUTPUT_FILE.read_text(encoding="utf-8")))
            store.set_setting("legacy_imported", "yes")
        defaults = {"device_ip": DEFAULT_DEVICE_IP, "port": str(DEFAULT_PORT),
                    "credentials": DEFAULT_CREDENTIALS_FILE, "spreadsheet_id": SPREADSHEET_ID,
                    "worksheet": WORKSHEET_NAME}
        AttendoSyncApp(root, store, fetch_dashboard_records, upload_to_google_sheets, defaults,
                      demo=args.demo, logo_path=find_resource(ICON_PNG))
        icon = find_resource(ICON_ICO)
        if icon:
            try:
                root.iconbitmap(str(icon))
            except Exception:
                pass
    except Exception as exc:
        messagebox.showerror("Attendo-Sync startup", str(exc), parent=root)
        root.destroy()
        return
    root.deiconify()
    root.mainloop()


if __name__ == "__main__":
    main()
