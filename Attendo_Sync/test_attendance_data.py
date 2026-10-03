import json
import os
import subprocess
import sys
import tempfile
import tkinter
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import Mock, patch

from attendance_data import AttendanceStore, daily_summary, monthly_summary


def punch(day, clock, mode, user_id="101"):
    return {"user_id": user_id, "user_name": "Anil", "timestamp": f"2026-09-{day}T{clock}:00",
            "status": 1, "punch": mode}


class AttendanceTests(unittest.TestCase):
    def test_first_scan_is_in_last_scan_is_out(self):
        daily = daily_summary([punch("01", "13:00", 0), punch("01", "09:00", 0),
                               punch("01", "12:00", 0), punch("01", "18:30", 0)])
        entry = daily[("101", date(2026, 9, 1))]
        self.assertEqual(entry["clock_in"].strftime("%H:%M"), "09:00")
        self.assertEqual(entry["clock_out"].strftime("%H:%M"), "18:30")
        self.assertEqual(entry["hours"], 9.5)
        self.assertTrue(entry["complete"])
        self.assertEqual(entry["issues"], "OK")

    def test_single_scan_is_missing_clock_out(self):
        entry = daily_summary([punch("01", "09:00", 0), punch("01", "09:00", 0),
                               {**punch("01", "09:00", 0), "timestamp": "2026-09-01T09:02:00"}])[("101", date(2026, 9, 1))]
        self.assertIsNone(entry["clock_out"])
        self.assertFalse(entry["complete"])
        self.assertEqual(entry["hours"], 0)
        self.assertEqual(entry["issues"], "Missing clock-out")

    def test_net_overtime_and_deficit_pay(self):
        employees = {"101": {"name": "Anil", "daily_pay": 100}}
        for clock_out, salary, net in [("17:30", 100, 0), ("18:30", 123.53, 1), ("16:30", 88.24, -1)]:
            with self.subTest(clock_out=clock_out):
                daily = daily_summary([punch("01", "09:00", 0), punch("01", clock_out, 1)])
                row = monthly_summary(employees, daily, "2026-09", date(2026, 10, 1))[0]
                self.assertEqual(row["salary"], salary)
                self.assertEqual(row["net_hours"], net)
                self.assertAlmostEqual(row["extra_days"], max(0, net / 8.5))

    def test_monthly_netting_and_incomplete_day(self):
        employees = {"101": {"name": "Anil", "daily_pay": 100}}
        rows = [punch("01", "09:00", 0), punch("01", "18:30", 1),
                punch("02", "09:00", 0), punch("02", "16:30", 1), punch("03", "09:00", 0)]
        summary = monthly_summary(employees, daily_summary(rows), "2026-09", date(2026, 10, 1))[0]
        self.assertEqual(summary["salary"], 200)
        self.assertEqual(summary["net_hours"], 0)
        self.assertEqual(summary["review_days"], 1)
        self.assertEqual(monthly_summary(employees, daily_summary(rows), "2026-08")[0]["salary"], 0)

    def test_punch_mode_is_ignored(self):
        daily = daily_summary([punch("01", "09:00", 4), punch("01", "18:00", 1)])
        self.assertTrue(daily[("101", date(2026, 9, 1))]["complete"])

    def test_admins_retired_and_employees_managed(self):
        with tempfile.TemporaryDirectory() as folder:
            store = AttendanceStore(Path(folder) / "attendance.db")
            store.merge([{**punch("28", "09:00", 0, "1"), "user_name": "Bijoy"},
                         {**punch("28", "17:30", 0, "1"), "user_name": "Bijoy"}],
                        {"1": "Bijoy", "2": "Aju", "3": "Treesa"})
            employees, rows = store.snapshot()
            self.assertEqual(list(employees), ["3"])
            self.assertEqual(rows, [])
            store.merge([], {"1": "Bijoy"})
            self.assertNotIn("1", store.snapshot()[0])
            store.add_employee("30", "  New   Person ", 120)
            self.assertEqual(store.snapshot()[0]["30"], {"name": "New Person", "daily_pay": 120})
            with self.assertRaises(ValueError):
                store.add_employee("30", "Duplicate")
            with self.assertRaises(ValueError):
                store.add_employee("31", "")
            store.remove_employee("30")
            store.merge([], {"30": "New Person"})
            self.assertNotIn("30", store.snapshot()[0])
            store.add_employee("30", "New Person")
            self.assertIn("30", store.snapshot()[0])

    def test_history_dedup_rates_and_atomic_merge(self):
        with tempfile.TemporaryDirectory() as folder:
            store = AttendanceStore(Path(folder) / "attendance.db")
            records = [punch("28", "09:00", 0), punch("28", "17:30", 1)]
            store.merge(records, {"102": "Bindu"})
            store.set_pay("101", 150)
            store.merge(records)
            store.merge([punch("29", "09:00", 0), punch("27", "09:00", 0)])
            employees, rows = store.snapshot()
            self.assertEqual(len(rows), 3)
            self.assertEqual(employees["101"]["daily_pay"], 150)
            self.assertIn("102", employees)
            with self.assertRaises(ValueError):
                store.merge([punch("30", "09:00", 0), {"user_id": "101"}])
            self.assertEqual(len(store.snapshot()[1]), 3)
            for amount in [-1, float("nan"), float("inf")]:
                with self.assertRaises(ValueError):
                    store.set_pay("101", amount)
            store.set_setting("device_ip", "192.168.29.201")
            self.assertEqual(AttendanceStore(store.path).setting("device_ip"), "192.168.29.201")


class IntegrationTests(unittest.TestCase):
    def test_backup_appends_missing_punches_without_clearing(self):
        import Attendo_Sync as app
        with tempfile.TemporaryDirectory() as folder:
            credentials = Path(folder) / "key.json"
            credentials.write_text(json.dumps({"type": "service_account", "client_email": "test@example.invalid",
                                               "private_key": "test-only", "token_uri": "test-only"}))
            old = punch("01", "09:00", 0)
            new = punch("01", "17:30", 1)
            worksheet = Mock()
            existing = app.build_worksheet_rows([old])
            worksheet.get_all_values.side_effect = lambda: existing
            worksheet.append_rows.side_effect = lambda records, **kwargs: existing.extend(records)
            client = Mock()
            client.open_by_key.return_value.worksheet.return_value = worksheet
            with patch.object(app.gspread, "service_account", return_value=client):
                app.upload_to_google_sheets([old, new, new], credentials, "Backup", log=lambda text: None)
                self.assertEqual(len(existing), 3)
                app.upload_to_google_sheets([old, new], credentials, "Backup", log=lambda text: None)
                self.assertEqual(worksheet.append_rows.call_count, 1)
                worksheet.clear.assert_not_called()
                existing[0] = ["incompatible"]
                with self.assertRaises(ValueError):
                    app.upload_to_google_sheets([old], credentials, "Backup", log=lambda text: None)

    def test_excel_values_and_untrusted_names(self):
        from attendance_reports import export_workbook
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as folder:
            rows = [punch("01", "09:00", 0), punch("01", "18:30", 1)]
            employees = {"101": {"name": "=1+1", "daily_pay": 100}, "102": {"name": "Bindu", "daily_pay": 100}}
            path = Path(folder) / "report.xlsx"
            export_workbook(path, employees, rows, "2026-09", today=date(2026, 9, 1))
            workbook = load_workbook(path)
            try:
                self.assertEqual(len(workbook.sheetnames), 5)
                self.assertEqual(workbook["Monthly payroll"]["I2"].value, 123.53)
                self.assertEqual(workbook["Monthly payroll"]["C2"].data_type, "s")
                self.assertEqual(workbook["Today"].max_row, 3)
                self.assertEqual(workbook["Raw punches"].max_row, 3)
            finally:
                workbook.close()

    def test_fetch_failures_preserve_local_data(self):
        from dashboard import AttendoSyncApp
        import queue
        with tempfile.TemporaryDirectory() as folder:
            store = AttendanceStore(Path(folder) / "history.sqlite3")
            records = [punch("28", "09:00", 0), punch("28", "17:30", 1)]
            worker = object.__new__(AttendoSyncApp)
            worker.store = store
            worker.events = queue.Queue()
            worker.fetch_attendance = Mock(return_value=(records, {"101": "Anil"}))
            worker.upload = Mock(side_effect=RuntimeError("Cloud unavailable"))
            settings = {"port": "4370", "credentials": "key", "worksheet": "Backup", "spreadsheet_id": "sheet"}
            worker._sync_worker("192.168.29.201", settings, True)
            self.assertEqual(len(store.snapshot()[1]), 2)
            messages = list(worker.events.queue)
            self.assertEqual(messages[-1][0], "error")
            self.assertIn("Local history saved", messages[-1][1])
            worker.fetch_attendance.side_effect = RuntimeError("Device offline")
            worker._sync_worker("192.168.29.201", settings, False)
            self.assertEqual(len(store.snapshot()[1]), 2)
            self.assertIn("existing history retained", list(worker.events.queue)[-1][1])
            worker.fetch_attendance.side_effect = None
            worker.upload.reset_mock()
            worker._sync_worker("192.168.29.201", settings, False)
            worker.upload.assert_not_called()
            self.assertEqual(len(store.snapshot()[1]), 2)

    def test_dashboard_at_high_dpi(self):
        for scaling in ("2.0", "2.6667", "3.3333"):
            with self.subTest(scaling=scaling):
                environment = os.environ.copy()
                environment["ATTENDO_TEST_TK_SCALING"] = scaling
                result = subprocess.run(
                    [sys.executable, "-m", "unittest",
                     "test_attendance_data.IntegrationTests.test_dashboard_tabs_filters_and_pay_edit", "-v"],
                    cwd=Path(__file__).resolve().parent, env=environment,
                    capture_output=True, text=True, timeout=90,
                )
                output = result.stdout + result.stderr
                self.assertEqual(result.returncode, 0, output)
                self.assertNotIn("axes sizes collapsed", output)

    def test_dashboard_tabs_filters_and_pay_edit(self):
        import Attendo_Sync as app
        from PIL import ImageGrab
        with tempfile.TemporaryDirectory() as folder:
            store = AttendanceStore(Path(folder) / "history.sqlite3")
            store.merge(app.load_sample_records())
            defaults = {"device_ip": app.DEFAULT_DEVICE_IP, "port": "4370", "credentials": "",
                        "spreadsheet_id": app.SPREADSHEET_ID, "worksheet": app.WORKSHEET_NAME}
            root = app.ttkbootstrap.Window(themename="flatly")
            if os.getenv("ATTENDO_TEST_TK_SCALING"):
                root.tk.call("tk", "scaling", float(os.environ["ATTENDO_TEST_TK_SCALING"]))
            view = app.AttendoSyncApp(root, store, Mock(), Mock(), defaults, logo_path=app.find_resource(app.ICON_PNG))
            try:
                self.assertEqual(len(view.notebook.tabs()), 3)
                self.assertEqual(len(view.daily_table.get_children()), 5)
                self.assertEqual(view.daily_table.cget("columns"), ("name", "in", "out", "hours", "status"))
                view.manage_employees()
                root.update()
                dialog = [child for child in root.winfo_children() if isinstance(child, tkinter.Toplevel)][-1]
                dialog.destroy()
                past = sorted(view.date_picker.cget("values"))[0]
                self.assertLess(past, date.today().isoformat())
                view.daily_date.set(past)
                view.refresh_daily()
                self.assertIn("OK", [view.daily_table.item(item)["values"][-1] for item in view.daily_table.get_children()])
                self.assertEqual(view.salary_value.cget("style"), "Metric.TLabel")
                self.assertGreaterEqual(int(app.ttkbootstrap.Style().lookup("Treeview", "rowheight")), 36)
                view.month.set("2026-09")
                view.refresh_month()
                names = [view.monthly_table.item(item)["values"][0] for item in view.monthly_table.get_children()]
                self.assertEqual(names, sorted(names, key=str.casefold))
                view.chart_period.set("Till date")
                view.draw_chart()
                self.assertEqual(len(view.chart_bars), len(view.employees))
                view.chart_period.set("2026-09")
                view.draw_chart()
                anil = [info for info in view.chart_info if info[0] == "Anil Kumar"][0]
                self.assertGreater(anil[1], 8)
                self.assertGreater(anil[2], 0)
                view.monthly_table.selection_set("101")
                with patch("dashboard.simpledialog.askfloat", return_value=150), patch("dashboard.messagebox.askokcancel", return_value=True):
                    view.edit_rate()
                self.assertEqual(store.snapshot()[0]["101"]["daily_pay"], 150)
                self.assertIn("150.00", view.monthly_table.item("101")["values"][-1])
                for width, height in [(1280, 860), (1000, 700), (1280, 860)]:
                    root.geometry(f"{width}x{height}+0+0")
                    for index in range(3):
                        view.notebook.select(index)
                        root.update()
                        view.canvas.draw()
                        self.assertTrue(view.progress.winfo_ismapped())
                        if index == 0:
                            self.assertTrue(view.navigation.winfo_ismapped())
                            self.assertGreater(view.navigation.winfo_height(), 20)
                            self.assertGreater(view.axes.get_window_extent().height, 110)
                            self.assertGreaterEqual(view.dashboard.winfo_height(), view.dashboard.winfo_reqheight())
                            view.dashboard_viewport.yview_moveto(1)
                            root.update()
                            self.assertAlmostEqual(view.dashboard_viewport.yview()[1], 1.0)
                            view.dashboard_viewport.yview_moveto(0)
                        for button in (view.fetch_btn, view.send_btn, view.export_btn, view.settings_btn):
                            self.assertTrue(button.winfo_ismapped())
                            self.assertGreater(button.winfo_width(), 50)
                            self.assertGreaterEqual(button.winfo_width(), button.winfo_reqwidth())
                            self.assertLessEqual(button.winfo_rootx() + button.winfo_width(), root.winfo_rootx() + root.winfo_width())
                        if os.getenv("ATTENDO_SCREENSHOT_DIR"):
                            output = Path(os.environ["ATTENDO_SCREENSHOT_DIR"])
                            output.mkdir(parents=True, exist_ok=True)
                            ImageGrab.grab(bbox=(root.winfo_rootx(), root.winfo_rooty(),
                                                root.winfo_rootx() + root.winfo_width(),
                                                root.winfo_rooty() + root.winfo_height())).save(output / f"tab-{index}-{width}.png")
            finally:
                view.close()


if __name__ == "__main__":
    unittest.main()